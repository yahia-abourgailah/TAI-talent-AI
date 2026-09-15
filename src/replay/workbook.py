"""Reads the legacy master workbook without changing it."""

import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

REQUIRED_COLUMNS: tuple[str, ...] = (
    "Name",
    "Age",
    "Title",
    "Employer",
    "Location",
    "Phone Number",
    "Email",
    "Profile URL",
    "Education",
    "Years Exp",
    "Platform",
    "Date Added",
    "Score",
    "Tier",
    "Recommendation",
    "Signals (Reasons to call)",
    "Flags (reasons of disqualification)",
)


class WorkbookError(Exception):
    """The workbook is missing, unreadable, or not shaped like the master record."""


@dataclass(frozen=True, slots=True)
class MasterRow:
    sheet_row: int
    values: dict[str, Any]


@dataclass(frozen=True, slots=True)
class MasterSheet:
    file_name: str
    sha256: str
    sheet_name: str
    sheet_names: tuple[str, ...]
    rows: tuple[MasterRow, ...]
    blank_rows_skipped: int
    columns: tuple[str, ...] = ()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _is_blank(cell: Any) -> bool:
    return cell is None or (isinstance(cell, str) and not cell.strip())


def read_master(path: Path, sheet: str | None = None) -> MasterSheet:
    """Every row with at least one non-blank cell. Rows that only carry formatting are skipped."""
    if not path.is_file():
        raise WorkbookError(f"No workbook at {path}. Set TALENT_MASTER_PATH or pass --master.")
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many unrelated types for a damaged file
        raise WorkbookError(f"Could not open {path.name}: {type(exc).__name__}") from exc

    try:
        if sheet is not None and sheet not in workbook.sheetnames:
            names = ", ".join(workbook.sheetnames)
            raise WorkbookError(f"No sheet named {sheet!r}. Sheets: {names}.")
        worksheet = workbook[sheet] if sheet is not None else workbook.worksheets[0]
        cells = worksheet.iter_rows(values_only=True)
        header = next(cells, None)
        if header is None:
            raise WorkbookError(f"Sheet {worksheet.title!r} is empty.")

        columns = tuple("" if name is None else str(name).strip() for name in header)
        missing = [name for name in REQUIRED_COLUMNS if name not in columns]
        if missing:
            raise WorkbookError(
                f"Sheet {worksheet.title!r} is missing columns: {', '.join(missing)}."
            )
        repeated = sorted({name for name in columns if name and columns.count(name) > 1})
        if repeated:
            raise WorkbookError(f"Columns appear more than once: {', '.join(repeated)}.")

        rows: list[MasterRow] = []
        blank = 0
        for sheet_row, values in enumerate(cells, start=2):
            if all(_is_blank(value) for value in values):
                blank += 1
                continue
            record = {
                name: (values[index] if index < len(values) else None)
                for index, name in enumerate(columns)
                if name
            }
            # A cell under a blank header, or past the last header, is kept, named by position.
            for index, value in enumerate(values):
                if (index >= len(columns) or not columns[index]) and not _is_blank(value):
                    record[f"(column {index + 1})"] = value
            rows.append(MasterRow(sheet_row, record))

        return MasterSheet(
            file_name=path.name,
            sha256=file_sha256(path),
            sheet_name=worksheet.title,
            sheet_names=tuple(workbook.sheetnames),
            rows=tuple(rows),
            blank_rows_skipped=blank,
            columns=tuple(name for name in columns if name),
        )
    finally:
        workbook.close()
