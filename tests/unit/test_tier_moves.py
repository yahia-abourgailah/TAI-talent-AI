"""BR-304: the who-moves-tier page, and the last-200 comparison, without a database."""

import textwrap
from datetime import date
from pathlib import Path

import openpyxl

from candidates.queue import reason_text
from replay.last_200 import load_legacy, most_recent, score_both
from replay.tier_moves import SAMPLE_SIZE, Move, render_html, sample
from replay.workbook import REQUIRED_COLUMNS, read_master
from scoring.rulesets import v2026_08_04 as ruleset

RULESET_FILE = Path(ruleset.__file__)


def _page(moves: list[Move], passed: bool = False) -> str:
    return render_html(
        title="Test",
        before_label="stored",
        after_label="replayed",
        facts=[("Run date", "2026-09-14")],
        moves=moves,
        verdict="verdict text",
        passed=passed,
    )


def test_the_sample_is_an_even_spread_of_moved_rows_only():
    moves = [Move(row, "P1", "P2") for row in range(2, 202)] + [Move(500, "P3", "P3")]
    picked = sample(moves)
    assert len(picked) == SAMPLE_SIZE
    assert all(m.before != m.after for m in picked)
    assert picked == sample(list(reversed(moves)))
    assert picked[0].sheet_row == 2 and picked[-1].sheet_row > 180


def test_the_page_counts_moves_by_direction():
    moves = [Move(2, "P1", "P2"), Move(3, "P1", "P2"), Move(4, "P3", "P2"), Move(5, "P4", "P4")]
    page = _page(moves)
    assert "<b>3</b><span>change tier" in page
    assert "<b>1</b><span>move to a better tier" in page
    assert "P1 → P2</td><td>2" in page
    assert "<script" not in page and "http" not in page


def test_a_page_with_no_moves_says_so():
    page = _page([Move(2, "P2", "P2")], passed=True)
    assert "Nobody changes tier." in page and 'class="verdict ok"' in page


def _master(tmp_path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for column, name in enumerate(REQUIRED_COLUMNS, start=1):
        sheet.cell(row=1, column=column, value=name)
    rows = [
        ("2026-09-01", "New Cairo"),
        ("2026-09-03", "Nasr City"),
        (None, "Maadi"),
        ("2026-09-02", "Sheikh Zayed"),
    ]
    for number, (added, place) in enumerate(rows, start=2):
        values = {
            "Name": f"Made-up {number}",
            "Title": "Sales Rep",
            "Location": place,
            "Age": 24,
            "Years Exp": 1,
            "Education": "bachelor",
            "Date Added": added,
        }
        for column, name in enumerate(REQUIRED_COLUMNS, start=1):
            if name in values:
                sheet.cell(row=number, column=column, value=values[name])
    path = tmp_path / "master.xlsx"
    workbook.save(path)
    return path


def test_the_most_recent_rows_come_first_and_undated_rows_last(tmp_path):
    sheet = read_master(_master(tmp_path))
    rows, undated = most_recent(sheet, 3)
    assert [r.sheet_row for r in rows] == [3, 5, 2]
    assert undated == 1


def test_the_ported_script_agrees_with_the_platform(tmp_path):
    sheet = read_master(_master(tmp_path))
    pairs = score_both(list(sheet.rows), load_legacy(RULESET_FILE), date(2026, 9, 17))
    assert pairs and not any(p.differs for p in pairs)


def test_a_changed_script_is_listed_with_the_part_that_moved(tmp_path):
    changed = tmp_path / "old_scorer.py"
    changed.write_text(
        RULESET_FILE.read_text(encoding="utf-8")
        + textwrap.dedent(
            """
            _real_location = _score_location
            def _score_location(c):
                points, flags, signals = _real_location(c)
                return (25 if points == 30 else points), flags, signals
            """
        ),
        encoding="utf-8",
    )
    sheet = read_master(_master(tmp_path))
    pairs = score_both(list(sheet.rows), load_legacy(changed), date(2026, 9, 17))
    moved = [p for p in pairs if p.differs]
    assert moved and all(p.cause == "location" for p in moved)
    assert all(p.legacy_score == p.platform_score - 5 for p in moved)


def test_every_review_line_reads_as_a_sentence_not_a_code():
    assert reason_text({"kind": "flagged_document", "reason_code": "ocr_timed_out"}).startswith(
        "We could not read this CV"
    )
    assert "Too far" in reason_text(
        {"kind": "proposed_rejection", "reason_code": "x", "rejection_label": "Too far"}
    )
    borderline = reason_text(
        {
            "kind": "borderline_score",
            "reason_code": "near_tier_line",
            "tier_above": "P2",
            "tier_below": "P3",
        }
    )
    assert "between P2 and P3" in borderline
    assert "_" not in reason_text({"kind": "flagged_document", "reason_code": "new_code"})
