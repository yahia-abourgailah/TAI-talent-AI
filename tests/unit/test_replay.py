"""The baseline replay tool, on a small fabricated workbook."""

import datetime
import json
import os
import stat
import subprocess
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from replay.baseline import main, run_baseline
from replay.clock import pinned_today
from replay.mapping import map_row
from replay.report import category, summarise
from replay.results import apply_rulings
from replay.rulings import KEEP_STORED, RULINGS, Ruling
from replay.workbook import REQUIRED_COLUMNS, WorkbookError, read_master
from scoring.rulesets.v2026_08_04 import Candidate, score_candidate

RUN_DATE = date(2026, 9, 14)
DAY = RUN_DATE.isoformat()
# Assembled so the repository's phone-number guard stays a check on real data.
FAKE_MOBILE = "010" + "00000000"
SENTINELS = (
    "Sentinel Zed Person",
    "sentinel@example.com",
    "https://example.com/in/sentinel-zed",
    FAKE_MOBILE,
)
COLUMNS = (*REQUIRED_COLUMNS, "Last Active", "Stage")


def _write(sheet, row_number: int, **values) -> None:
    for column_number, name in enumerate(COLUMNS, start=1):
        if name in values:
            sheet.cell(row=row_number, column=column_number, value=values[name])


def _scored_by_the_ruleset() -> dict[str, object]:
    values: dict[str, object] = {
        "Name": SENTINELS[0],
        "Email": SENTINELS[1],
        "Profile URL": SENTINELS[2],
        "Phone Number": FAKE_MOBILE,
        "Title": "Sales Rep",
        "Location": "New Cairo",
        "Years Exp": 1,
        "Education": "bachelor",
        "Age": "25",
        "Last Active": "2026-09-10",
    }
    candidate = Candidate(
        full_name=SENTINELS[0],
        email=SENTINELS[1],
        profile_url=SENTINELS[2],
        phone_number=FAKE_MOBILE,
        current_title="Sales Rep",
        location="New Cairo",
        years_experience=1.0,
        education_level="bachelor",
        age=25,
    )
    with pinned_today(RUN_DATE):
        result = score_candidate(candidate)
    values.update(
        {
            "Score": result.overall_score,
            "Tier": result.priority,
            "Recommendation": result.recommendation,
            "Signals (Reasons to call)": "; ".join(result.key_signals),
            "Flags (reasons of disqualification)": "; ".join(result.red_flags),
        }
    )
    return values


@pytest.fixture
def master(tmp_path: Path) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for column_number, name in enumerate(COLUMNS, start=1):
        sheet.cell(row=1, column=column_number, value=name)
    _write(sheet, 2, **_scored_by_the_ruleset())
    _write(
        sheet,
        3,
        Name="Mismatch",
        Title="Sales, customer service, retail and hospitality",  # several signals
        # Two private-university names match; the ruleset reports whichever its set yields first.
        Education="bachelor, American University in Cairo (AUC)",
        Location="Maadi",
        Score=99,
        Tier="P1",
    )
    sheet.cell(row=4, column=1).number_format = "0"  # formatting only, no value
    sheet.cell(row=5, column=1, value="   ")  # whitespace only
    _write(
        sheet, 6, **{"Name": "Never", "Title": "Customer Service", "Age": "?", "Years Exp": "3+"}
    )
    path = tmp_path / "data" / "TAI_Master.xlsx"
    path.parent.mkdir()
    workbook.save(path)
    return path


def _run(master: Path, out: Path, *extra: str) -> int:
    return main(["--master", str(master), "--out", str(out), "--run-date", DAY, *extra])


def test_reads_non_blank_rows_and_keeps_sheet_row_numbers(master):
    sheet = read_master(master)
    assert [row.sheet_row for row in sheet.rows] == [2, 3, 6]
    assert sheet.blank_rows_skipped == 2
    assert len(sheet.sha256) == 64


def test_missing_columns_are_named(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.append(["Name", "Score"])
    path = tmp_path / "partial.xlsx"
    workbook.save(path)
    with pytest.raises(WorkbookError, match="missing columns: Age"):
        read_master(path)


def test_unknown_markers_stay_empty_and_unreadable_numbers_are_recorded(master):
    mapped = map_row(read_master(master).rows[2])
    assert mapped.candidate.age is None
    assert mapped.candidate.years_experience is None
    assert mapped.parse_issues == ("Years Exp",)


def test_last_active_is_not_given_to_the_ruleset(master):
    """The stored scores were produced without it; see replay/mapping.py."""
    row = read_master(master).rows[0]
    assert row.values["Last Active"] == "2026-09-10"
    assert map_row(row).candidate.last_active == ""


def test_replay_matches_a_row_scored_by_the_same_rules(master):
    matching, differing, never_scored = run_baseline(read_master(master), RUN_DATE)
    assert matching.score_match and matching.tier_match
    assert not differing.score_match
    assert not never_scored.has_stored_score
    assert never_scored.replayed.tier.startswith("P")


def test_run_date_is_pinned_and_then_released():
    candidate = Candidate(location="New Cairo", last_active="2026-09-10")
    with pinned_today(date(2026, 9, 12)):
        recent = score_candidate(candidate)
    with pinned_today(date(2027, 1, 1)):
        stale = score_candidate(candidate)
    assert any("Active recently" in signal for signal in recent.key_signals)
    assert any(flag.startswith("Profile stale") for flag in stale.red_flags)
    assert datetime.date is date


def test_category_keeps_the_wording_and_drops_the_candidate():
    assert category("Near New Cairo: Fifth Settlement") == "Near New Cairo"
    assert (
        category("DISQUALIFIED: Managerial/director title: 'Head'") == "Managerial/director title"
    )
    assert category("3 years — mid entry level") == "# years — mid entry level"
    assert category("★ Recent grad (2025) — priority") == "Recent grad"


def test_same_inputs_give_byte_identical_files(master, tmp_path):
    assert _run(master, tmp_path / "a") == 0
    assert _run(master, tmp_path / "b") == 0
    for name in (f"baseline-{DAY}.jsonl", f"baseline-{DAY}.json", f"report-{DAY}.md"):
        assert (tmp_path / "a" / name).read_bytes() == (tmp_path / "b" / name).read_bytes()


def test_files_do_not_depend_on_the_python_process(master, tmp_path):
    """The ruleset iterates over sets, whose order changes with each process's hash seed."""
    source = Path(__file__).resolve().parents[2] / "src"
    for seed in ("1", "2", "3"):
        subprocess.run(
            [
                sys.executable,
                "-m",
                "replay.baseline",
                *("--master", str(master), "--out", str(tmp_path / seed), "--run-date", DAY),
            ],
            check=True,
            capture_output=True,
            env={**os.environ, "PYTHONHASHSEED": seed, "PYTHONPATH": str(source)},
        )
    for name in (f"baseline-{DAY}.jsonl", f"baseline-{DAY}.json", f"report-{DAY}.md"):
        first = (tmp_path / "1" / name).read_bytes()
        assert first == (tmp_path / "2" / name).read_bytes() == (tmp_path / "3" / name).read_bytes()


def test_summary_counts(master, tmp_path):
    assert _run(master, tmp_path / "out") == 0
    summary = json.loads((tmp_path / "out" / f"baseline-{DAY}.json").read_text())["summary"]
    assert (summary["rows"], summary["stored_scores"], summary["score_matches"]) == (3, 2, 1)
    assert summary["unreadable_numbers"] == {"Years Exp": 1}
    assert (summary["unexplained_differences"], summary["parity"]) == (1, False)


def test_row_file_is_owner_only_and_the_report_holds_no_candidate_values(master, tmp_path, capsys):
    copy = tmp_path / "copy" / "BASELINE_REPORT.md"
    assert _run(master, tmp_path / "out", "--report-copy", str(copy)) == 0

    rows = tmp_path / "out" / f"baseline-{DAY}.jsonl"
    assert stat.S_IMODE(rows.stat().st_mode) == 0o600
    printed = capsys.readouterr().out
    for text in ((tmp_path / "out" / f"report-{DAY}.md").read_text(), copy.read_text(), printed):
        for sentinel in SENTINELS:
            assert sentinel not in text


def test_refuses_to_write_rows_inside_a_git_repository(master, tmp_path):
    repository = tmp_path / "repository"
    (repository / ".git").mkdir(parents=True)
    assert _run(master, repository / "baseline") == 2
    assert not (repository / "baseline").exists()


def test_refuses_to_read_candidate_data_from_inside_a_git_repository(master, tmp_path):
    repository = tmp_path / "repository"
    (repository / ".git").mkdir(parents=True)
    inside = repository / "TAI_Master.xlsx"
    inside.write_bytes(master.read_bytes())
    assert _run(inside, tmp_path / "out") == 2


def test_missing_workbook_is_a_clear_error(tmp_path, capsys):
    assert _run(tmp_path / "missing.xlsx", tmp_path / "out") == 2
    assert "No workbook" in capsys.readouterr().err


def _ruling_for_the_differing_row(sheet, **changes) -> Ruling:
    differing = run_baseline(sheet, RUN_DATE)[1]
    values = {
        "workbook_sha256": sheet.sha256,
        "sheet_row": differing.sheet_row,
        "stored_score": differing.stored.score,
        "stored_tier": differing.stored.tier,
        "replayed_score": differing.replayed.score,
        "replayed_tier": differing.replayed.tier,
        "decision": KEEP_STORED,
        "reason": "Fabricated ruling for a test",
        "recorded_on": DAY,
    }
    values.update(changes)
    return Ruling(**values)


def _summary_with(sheet, rulings):
    results, stale = apply_rulings(run_baseline(sheet, RUN_DATE), sheet.sha256, rulings)
    return summarise(results, [], stale)


def test_a_matching_ruling_explains_the_difference(master):
    sheet = read_master(master)
    summary = _summary_with(sheet, [_ruling_for_the_differing_row(sheet)])
    assert (summary["ruled_differences"], summary["unexplained_differences"]) == (1, 0)
    assert summary["parity"]
    assert summary["stale_rulings"] == []


def test_a_ruling_stops_applying_when_the_outcome_changes(master):
    sheet = read_master(master)
    summary = _summary_with(sheet, [_ruling_for_the_differing_row(sheet, replayed_score=-1)])
    assert (summary["unexplained_differences"], summary["parity"]) == (1, False)
    assert [stale["sheet_row"] for stale in summary["stale_rulings"]] == [3]


def test_a_ruling_for_another_workbook_is_ignored(master):
    sheet = read_master(master)
    other = _ruling_for_the_differing_row(sheet, workbook_sha256="0" * 64)
    summary = _summary_with(sheet, [other])
    assert (summary["unexplained_differences"], summary["stale_rulings"]) == (1, [])


def test_two_rulings_for_one_row_are_refused(master):
    sheet = read_master(master)
    ruling = _ruling_for_the_differing_row(sheet)
    with pytest.raises(ValueError, match="sheet row 3"):
        apply_rulings(run_baseline(sheet, RUN_DATE), sheet.sha256, [ruling, ruling])


def test_require_parity_fails_until_every_difference_is_ruled(master, tmp_path, monkeypatch):
    assert _run(master, tmp_path / "before", "--require-parity") == 1

    ruling = _ruling_for_the_differing_row(read_master(master))
    monkeypatch.setattr("replay.baseline.RULINGS", (ruling,))
    assert _run(master, tmp_path / "after", "--require-parity") == 0
    report = (tmp_path / "after" / f"report-{DAY}.md").read_text()
    assert "Parity reached" in report
    assert "Fabricated ruling for a test" in report


def test_recorded_rulings_name_rows_not_people():
    for ruling in RULINGS:
        assert len(ruling.workbook_sha256) == 64
        assert ruling.sheet_row > 1
        assert "@" not in ruling.reason
        assert "http" not in ruling.reason


def test_an_age_of_0_is_not_an_age_but_0_years_experience_is_kept():
    """A source that writes 0 for an unknown age must not disqualify the candidate as under 21."""
    from replay.workbook import MasterRow

    mapped = map_row(MasterRow(sheet_row=2, values={"Age": "0", "Years Exp": "0"}))
    assert mapped.candidate.age is None
    assert mapped.candidate.years_experience == 0.0
    assert map_row(MasterRow(sheet_row=3, values={"Age": 0})).candidate.age is None
