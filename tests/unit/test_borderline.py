"""BR-310 without a database: the tier lines match the ruleset, and each rule picks the right
scores. Made-up profiles only."""

import inspect
import re
from datetime import date

import openpyxl
import pytest

from replay.borderline_options import FLIPS, render, score_sheet
from replay.workbook import REQUIRED_COLUMNS, read_master
from scoring.borderline import LINES, Policy, near_line
from scoring.rulesets import v2026_08_04 as ruleset
from scoring.rulesets.v2026_08_04 import Candidate

VERSION = "2026-08-04"


@pytest.mark.parametrize("line,above,below", LINES[VERSION]["entry"])
def test_track_a_lines_are_the_rulesets_own(line, above, below):
    assert ruleset._recommend(line)[1] == above
    assert ruleset._recommend(line - 1)[1] == below


def test_track_b_lines_are_the_rulesets_own():
    source = inspect.getsource(ruleset._score_headhunt)
    found = [int(n) for n in re.findall(r"overall_score >= (\d+)", source)]
    assert tuple(found) == tuple(line for line, _, _ in LINES[VERSION]["headhunt"])


@pytest.mark.parametrize(
    "score,expected",
    [
        (53, None),
        (54, (55, "P2", "P3", -1)),
        (55, (55, "P2", "P3", 0)),
        (56, (55, "P2", "P3", 1)),
        (57, None),
        (74, (75, "P1", "P2", -1)),
        (34, (35, "P3", "P4", -1)),
        (0, None),
        (100, None),
    ],
)
def test_a_band_of_one_includes_the_line_and_both_sides(score, expected):
    found = near_line(score, VERSION, "entry", Policy("band", 1))
    got = (
        None if found is None else (found.line, found.tier_above, found.tier_below, found.distance)
    )
    assert got == expected


def test_below_the_line_takes_only_the_near_misses():
    policy = Policy("below_line", 2)
    assert near_line(53, VERSION, "entry", policy).distance == -2  # type: ignore[union-attr]
    assert near_line(55, VERSION, "entry", policy) is None
    assert near_line(56, VERSION, "entry", policy) is None


def test_track_b_uses_its_own_lines():
    found = near_line(79, VERSION, "headhunt", Policy("band", 1))
    assert found is not None and (found.tier_above, found.tier_below) == ("T1", "T2")
    assert near_line(75, VERSION, "headhunt", Policy("band", 1)) is None


@pytest.mark.parametrize("rule,points", [("band", 0), ("band", 10), ("near", 1)])
def test_a_rule_outside_the_agreed_shapes_is_refused(rule, points):
    with pytest.raises(ValueError):
        Policy(rule, points)  # type: ignore[arg-type]


def test_an_unknown_version_has_no_lines():
    with pytest.raises(KeyError):
        near_line(50, "2099-01-01", "entry", Policy("band", 1))


def test_the_explanation_names_the_line_the_tiers_and_the_rule():
    policy = Policy("band", 2)
    text = near_line(53, VERSION, "entry", policy).explain(policy, VERSION)  # type: ignore[union-attr]
    assert text == (
        "BORDERLINE: 2 points under the P2 line (55); between P2 and P3. "
        "Rule: within 2 of a tier line, criteria 2026-08-04"
    )


def test_the_age_flips_move_one_year_and_skip_a_candidate_without_an_age():
    person = Candidate(age=25, location="New Cairo")
    assert FLIPS["age one year younger"](person).age == 24  # type: ignore[union-attr]
    assert FLIPS["age one year older"](person).age == 26  # type: ignore[union-attr]
    assert FLIPS["age one year younger"](Candidate(location="New Cairo")) is None


def _place_worth(points: int) -> str:
    places = ruleset.NEAR_NEW_CAIRO if points == 30 else ruleset.OTHER_CAIRO
    return next(
        p for p in sorted(places) if ruleset._score_location(Candidate(location=p))[0] == points
    )


@pytest.mark.parametrize("points,flipped_points", [(30, 15), (15, 30)])
def test_the_location_flip_reads_the_other_cairo_band(points, flipped_points):
    flip = FLIPS["location read as the other Cairo band"]
    flipped = flip(Candidate(location=_place_worth(points)))
    assert flipped is not None
    assert ruleset._score_location(flipped)[0] == flipped_points
    assert not ruleset._check_hard_disqualifiers(flipped)[0]
    assert flip(Candidate(location="")) is None


def test_the_options_report_counts_without_candidate_values(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for column, name in enumerate(REQUIRED_COLUMNS, start=1):
        sheet.cell(row=1, column=column, value=name)
    rows = [
        {"Name": "Sentinel Qx One", "Title": "Sales Rep", "Location": "New Cairo", "Age": 21},
        {"Name": "Sentinel Qx Two", "Title": "Cashier", "Location": "Nasr City", "Age": 30},
    ]
    for number, values in enumerate(rows, start=2):
        for column, name in enumerate(REQUIRED_COLUMNS, start=1):
            if name in values:
                sheet.cell(row=number, column=column, value=values[name])
    path = tmp_path / "master.xlsx"
    workbook.save(path)

    master = read_master(path)
    run_date = date(2026, 9, 14)
    report = render(master, run_date, score_sheet(master, run_date))
    assert "| ±1 |" in report and "Option 3" in report
    assert "Sentinel Qx" not in report and "New Cairo" not in report
